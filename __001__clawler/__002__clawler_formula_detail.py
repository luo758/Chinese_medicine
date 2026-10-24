"""按方剂列表爬取详情，每个方剂保存为一个 txt。"""
import os
import sys
import time

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from tqdm import tqdm

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "formula_list.xlsx")
OUT_DIR = os.path.join(BASE_DIR, "方剂")
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    )
}


def fetch_text(session, url):
    resp = session.get(url, headers=HEADERS, timeout=20)
    resp.raise_for_status()
    resp.encoding = "utf-8"
    soup = BeautifulSoup(resp.text, "html.parser")
    content = soup.select_one("#content .p_content")
    if not content:
        return ""
    lines = []
    for tag in content.find_all(["h2", "h3", "p", "li"]):
        text = " ".join(tag.get_text(" ", strip=True).split())
        if text:
            lines.append(text)
    return "\n".join(lines)


def crawl_formula_detail(resume=True):
    # 创建输出目录
    os.makedirs(OUT_DIR, exist_ok=True)
    ws = load_workbook(EXCEL_FILE).active
    rows = [(name, url) for name, url in ws.iter_rows(min_row=2, values_only=True) if name and url]
    session = requests.Session()
    ok, failed = 0, []
    bar = tqdm(rows, desc="爬取方剂")
    for name, url in bar:
        bar.set_postfix_str(name, refresh=False)
        # 创建输出文件路径
        path = os.path.join(OUT_DIR, f"{name}.txt")
        if resume and os.path.exists(path) and os.path.getsize(path) > 0:
            # 跳过已存在的文件
            ok += 1
            continue
        try:
            text = fetch_text(session, url)
        except Exception as e:
            failed.append(name)
            tqdm.write(f"失败 {name}: {e}")
            continue
        if not text:
            failed.append(name)
            tqdm.write(f"空内容: {name}")
            continue
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
        ok += 1
        time.sleep(0.2)
    print(f"完成 {ok} 条，失败 {len(failed)} 条，目录 {OUT_DIR}")


if __name__ == "__main__":
    crawl_formula_detail(resume=False)
